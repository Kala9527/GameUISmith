"""Regenerate the Windows local deployment guide."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("Windows本地部署指南.docx")


def set_font(style, size, bold=False):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), "Microsoft YaHei")


def set_cell_border(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    tcpr.append(element)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
set_font(styles["Normal"], 10.5)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.2
set_font(styles["Title"], 18, True)
styles["Title"].paragraph_format.space_after = Pt(13)
set_font(styles["Heading 1"], 12, True)
styles["Heading 1"].paragraph_format.space_before = Pt(13)
styles["Heading 1"].paragraph_format.space_after = Pt(6)

doc.add_paragraph("游戏 UI 素材制作台 Windows 本地部署指南", "Title")
doc.add_paragraph(
    "本指南供收到项目文件夹的 Windows 用户使用。将整个项目文件夹放在可写位置，双击根目录的“启动项目.bat”即可检查环境、首次安装依赖并打开浏览器。生成类功能还需要用户自行提供可用的模型服务地址和 API Key。"
)

doc.add_paragraph("按 Node 环境操作", "Heading 1")
table = doc.add_table(rows=1, cols=3)
table.autofit = False
for cell, text, width in zip(
    table.rows[0].cells,
    ("当前情况", "处理方式", "启动结果"),
    (1.45, 3.55, 1.8),
):
    cell.width = Inches(width)
    cell.text = text
    shade(cell, "E6E6E6")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)
    cell.paragraphs[0].runs[0].bold = True

rows = [
    (
        "已安装符合要求的 Node.js",
        "直接双击“启动项目.bat”；脚本检查 npm，并在缺少依赖时执行 npm ci。",
        "自动启动并打开本机浏览器。",
    ),
    (
        "Node.js 版本过低",
        "按脚本提示从 Node.js 官网安装当前 LTS 版本。安装后重新打开命令窗口，再次双击脚本。",
        "不会在旧版本下继续安装依赖。",
    ),
    (
        "没有 Node.js",
        "从 https://nodejs.org/en/download 获取官方 Windows 安装程序，安装时保留 npm 和加入 PATH；然后重新双击脚本。",
        "安装 Node.js 后再启动。",
    ),
]
for index, values in enumerate(rows):
    row = table.add_row()
    for cell, value in zip(row.cells, values):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
        if index % 2:
            shade(cell, "F7F7F7")
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.space_before = Pt(3)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)

doc.add_paragraph(
    "最低兼容版本：Node.js 20.19.0（20 系列），或 Node.js 22.12.0 及更新的 22 系列，以及 23 系列及以上。建议安装官网当前 LTS，而不是为了满足最低值专门安装旧版本。"
)

doc.add_paragraph("首次启动与日常使用", "Heading 1")
for text in (
    "1. 保持项目目录完整，特别是 package.json、package-lock.json、src 和“启动项目.bat”。首次安装依赖需要网络能访问 npm 仓库，并且项目目录可写。",
    "2. 双击“启动项目.bat”。若 node_modules 中尚无 Vite，脚本执行 npm ci；安装失败会保留窗口显示错误，不会假装启动成功。后续启动若依赖已在，则直接运行。",
    "3. Vite 仅监听 127.0.0.1，并尝试自动打开浏览器。若没有自动打开，使用命令窗口中显示的 Local 地址。使用期间保持命令窗口打开；按 Ctrl+C 关闭服务。",
):
    doc.add_paragraph(text)

doc.add_paragraph("配置与使用范围", "Heading 1")
doc.add_paragraph(
    "打开应用的“模型 API 配置”，填写服务 Base URL、API Key 和对应任务的模型名称，并先测试提示词模型。模型服务和授权由使用者提供；网页能打开不代表 AI 接口一定可用。密钥会参与浏览器与本地服务间的请求，不要在共享电脑或不可信网络环境中使用他人的密钥。"
)
doc.add_paragraph(
    "该启动方式是本机工具，不是公网生产服务器。项目的跨域回退中转挂在 Vite 开发服务器上；不要把 --host 改为 0.0.0.0，也不要将端口转发到公网。仅打开 dist 文件或改用 vite preview，可能让需要中转的模型请求失败。"
)

doc.add_paragraph("常见问题", "Heading 1")
for lead, rest in (
    ("找不到 node 或 npm：", "确认已安装官方 Node.js 并重新打开命令窗口；在命令提示符运行 node -v 和 npm -v。"),
    ("npm ci 失败：", "检查联网、代理、npm 仓库访问和目录写入权限；不要用 npm audit fix 代替安装。"),
    ("安装后依赖异常：", "关闭服务，删除项目内的 node_modules 文件夹，再运行脚本以重新执行 npm ci。不要删除 package-lock.json。"),
    ("安全告警：", "npm 的漏洞告警不等于启动失败。交付前应审查 npm audit 结果、修复影响实际使用的漏洞并复测；脚本不会自动修改依赖版本。"),
):
    paragraph = doc.add_paragraph()
    paragraph.add_run(lead).bold = True
    paragraph.add_run(rest)

doc.save(OUTPUT)
print(OUTPUT)

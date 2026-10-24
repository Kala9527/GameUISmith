"""Merge the local deployment guide with the original illustrated usage guide."""

from io import BytesIO
from pathlib import Path
import sys

from docx import Document
from docx.shared import Inches


docs = Path(sys.argv[1]).resolve()
deployment = docs / "Windows本地部署指南.docx"
original = docs / "使用工作台配置及简单的调用方法 原稿.docx"
output = docs / "使用工作台配置及简单的调用方法.docx"

doc = Document(deployment)
source = Document(original)
images = {
    str(part.partname).split("/")[-1]: part.blob
    for part in source.part.related_parts.values()
    if str(part.partname).startswith("/word/media/")
}
if set(images) != {"image1.png", "image2.png", "image3.png"}:
    raise ValueError("The original guide's three screenshots were not found")

doc.paragraphs[0].text = "游戏 UI 素材制作台本地部署与生图使用指南"
doc.paragraphs[1].text = (
    "本指南供收到项目文件夹的 Windows 用户使用。先双击“启动项目.bat”检查环境并启动本机网页，"
    "再配置自己的模型服务，按截图完成提示词编写、图片生成与下载。生成类功能需要可用的模型服务地址和 API Key。"
)


def add_step(title, text, image_name, width):
    doc.add_paragraph(title, "Heading 1")
    doc.add_paragraph(text)
    image = doc.add_paragraph()
    image.paragraph_format.keep_together = True
    image.alignment = 1
    image.add_run().add_picture(BytesIO(images[image_name]), width=Inches(width))


add_step(
    "第一步 配置模型服务",
    "在“模型配置”中填写服务 Base URL、API Key 和对应任务的模型名称。示例截图中的地址仅是演示，"
    "请替换成自己实际获授权的服务地址；应用会自动补 /v1。确认密钥有目标模型的调用权限，"
    "点击“测试提示词生成模型”检查连通性。不要把密钥发给他人。",
    "image1.png",
    4.25,
)
add_step(
    "第二步 编写并套用提示词",
    "进入“提示词生成”，选择目标类型，描述想要的主体、用途、风格、背景、尺寸和不希望出现的元素。"
    "可以点击“让 AI 生成提示词配方”，也可以直接手动填写并修改正向、风格、反向提示。"
    "检查组合预览后点击“套用并去生图”；如图中所示，提示词会进入相应生成界面。",
    "image2.png",
    6.75,
)
add_step(
    "第三步 生成并保存图片",
    "在“文生图”界面确认提示词、尺寸、质量及模型，点击“调用图像模型生成”。"
    "结果出现后在“生成结果”中点击“下载”保存到本机；不要只依赖当前网页中的预览。"
    "如果生成失败，先核对 Key、模型权限和服务地址，再查看窗口或页面中的错误提示。",
    "image3.png",
    6.75,
)

doc.add_paragraph(
    "补充资料：同目录的“图片生成API调用文档.md”和 TXT 版本说明接口参数与直接调用方法。"
    "这些 API 示例与工作台界面使用方式不同，第三方网关的具体返回格式和模型支持应向服务提供方确认。"
)
doc.save(output)
print(output)
